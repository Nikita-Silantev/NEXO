# -*- coding: utf-8 -*-
"""Генератор документов .docx по оформлению ЕСПД (ГОСТ 19.104, 19.105, 19.106).
Двухпроходная сборка: после первого прохода документ переводится в PDF, определяются страницы
заголовков и общее число листов, затем документ собирается повторно с заполненным содержанием."""
import re, os, subprocess, tempfile, shutil
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

FONT = "Times New Roman"
MONO = "Courier New"
TEXT_W = {"A4": 17.0, "A4L": 25.7, "A3L": 38.0, "A3P": 25.7}
PAGE = {"A4": (21.0, 29.7), "A4L": (29.7, 21.0), "A3L": (42.0, 29.7), "A3P": (29.7, 42.0)}  # (ширина, высота) в альбомной - уже перевёрнуто


# ---------------------------------------------------------------- низкоуровневые помощники
def _set_font(style, name=FONT, size=14, bold=None, italic=None, caps=None):
    f = style.font
    f.name = name
    f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if caps is not None: f.all_caps = caps
    f.color.rgb = None
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    col = rpr.find(qn("w:color"))
    if col is not None: rpr.remove(col)


def _run_font(run, name=FONT, size=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic


def _fld(par, instr, size=12, text="1"):
    def mk(t):
        r = par.add_run(); _run_font(r, size=size)
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e); return r
    mk("begin")
    r = par.add_run(); _run_font(r, size=size)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    mk("separate")
    r = par.add_run(text); _run_font(r, size=size)
    mk("end")


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def _cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def _borders(table, sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e)
    tblPr.append(b)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|__.+?__)")


class Doc:
    def __init__(self, designation, title, program="Мессенджер NEXO", kind="espd", doc_code="ТЗ 01",
                 org="Силантьев Никита (владелец проекта)", year="2026", version="1.0", subtitle=None,
                 approval_label="УТВЕРЖДАЮ"):
        self.designation = designation
        self.title = title
        self.program = program
        self.kind = kind  # 'espd' | 'plain'
        self.doc_code = doc_code
        self.org = org
        self.year = year
        self.version = version
        self.subtitle = subtitle
        self.blocks = []
        self.keys = {}
        self.fig_n = 0
        self.tab_n = 0
        self.hn = [0, 0, 0]
        self.app_n = 0
        self.toc_entries = []  # (level, label)
        self.want_toc = kind == "espd"
        self.refkey = None
        self.ext_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "build")

    # ------------------------------------------------------------ API авторов
    def h1(self, text, key=None, numbered=True):
        if numbered:
            self.hn = [self.hn[0] + 1, 0, 0]
            label = f"{self.hn[0]} {text}"
            num = str(self.hn[0])
        else:
            label, num = text, ""
        if key: self.keys["sec:" + key] = num
        self.toc_entries.append((1, label))
        self.blocks.append(dict(t="h", lvl=1, label=label, text=text, num=num))

    def h2(self, text, key=None):
        self.hn = [self.hn[0], self.hn[1] + 1, 0]
        num = f"{self.hn[0]}.{self.hn[1]}"
        label = f"{num} {text}"
        if key: self.keys["sec:" + key] = num
        self.toc_entries.append((2, label))
        self.blocks.append(dict(t="h", lvl=2, label=label, text=text, num=num))

    def h3(self, text, key=None):
        self.hn = [self.hn[0], self.hn[1], self.hn[2] + 1]
        num = f"{self.hn[0]}.{self.hn[1]}.{self.hn[2]}"
        label = f"{num} {text}"
        if key: self.keys["sec:" + key] = num
        self.blocks.append(dict(t="h", lvl=3, label=label, text=text, num=num))

    def appendix(self, title, status="справочное", key=None):
        self.app_n += 1
        label = f"ПРИЛОЖЕНИЕ {self.app_n} {title}"
        if key: self.keys["app:" + key] = str(self.app_n)
        self.toc_entries.append((1, label))
        self.blocks.append(dict(t="app", n=self.app_n, status=status, title=title, label=label))

    def p(self, text, indent=True, align="j", bold=False, size=None, keep=False):
        self.blocks.append(dict(t="p", text=text, indent=indent, align=align, bold=bold, size=size, keep=keep))

    def bullets(self, items, mark="–"):
        for it in items:
            self.blocks.append(dict(t="li", text=it, mark=mark))

    def numbered(self, items, style="a"):
        letters = "абвгдежиклмнпрстуфхцчшэюя"
        for i, it in enumerate(items):
            m = f"{letters[i]})" if style == "a" else f"{i + 1})"
            self.blocks.append(dict(t="li", text=it, mark=m))

    def note(self, text, label="Примечание"):
        self.blocks.append(dict(t="note", text=text, label=label))

    def code(self, text, size=9):
        self.blocks.append(dict(t="code", text=text.strip("\n"), size=size))

    def table(self, caption, header, rows, widths=None, key=None, font=12, sect=None, nocaption=False):
        if not nocaption:
            self.tab_n += 1
            n = self.tab_n
            if key: self.keys["tab:" + key] = str(n)
        else:
            n = None
        self.blocks.append(dict(t="tab", n=n, caption=caption, header=header, rows=rows, widths=widths, font=font, sect=sect))

    def fig(self, path, caption, key=None, dpi=130, big=None):
        self.fig_n += 1
        if key: self.keys["fig:" + key] = str(self.fig_n)
        w, h = Image.open(path).size
        nw, nh = w / dpi * 2.54, h / dpi * 2.54
        s_a4 = min(17.0 / nw, 20.5 / nh, 1.0)
        s_l = min(38.0 / nw, 22.0 / nh, 1.0)
        s_p = min(25.7 / nw, 35.5 / nh, 1.0)
        sect = None
        if big is True or (big is None and s_a4 < 0.55):
            sect = "A3L" if s_l >= s_p else "A3P"
        self.blocks.append(dict(t="fig", n=self.fig_n, path=path, caption=caption, nw=nw, nh=nh, big=bool(sect), sect=sect))

    def landscape(self):
        self.blocks.append(dict(t="mode", sect="A4L"))

    def portrait(self):
        self.blocks.append(dict(t="mode", sect=None))

    def pagebreak(self):
        self.blocks.append(dict(t="pb"))

    def spacer(self, pts=6):
        self.blocks.append(dict(t="sp", pts=pts))

    # ------------------------------------------------------------ разрешение ссылок
    def _res(self, s):
        def rep(m):
            k = m.group(1)
            if k not in self.keys:
                raise KeyError(f"Неизвестная ссылка {{{k}}} в документе {self.designation}")
            return self.keys[k]
        def rep_ext(m):
            dk, k = m.group(1), m.group(2)
            f = os.path.join(self.ext_dir, f"refs_{dk}.json")
            if not os.path.exists(f):
                raise KeyError(f"Нет файла ссылок {f}: соберите документ {dk} раньше")
            import json
            ks = json.load(open(f, encoding="utf-8"))
            if k not in ks: raise KeyError(f"Неизвестная внешняя ссылка {dk}:{k}")
            return ks[k]
        s = re.sub(r"\{(\w+):((?:fig|tab|sec|app):[\w\-]+)\}", rep_ext, s)
        return re.sub(r"\{((?:fig|tab|sec|app):[\w\-]+)\}", rep, s)

    # ------------------------------------------------------------ построение документа
    def _new_doc(self):
        d = Document()
        st = d.styles
        _set_font(st["Normal"], size=14)
        pf = st["Normal"].paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.5
        for name, size, caps, align in (("Heading 1", 14, True, WD_ALIGN_PARAGRAPH.CENTER),
                                        ("Heading 2", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
                                        ("Heading 3", 14, False, WD_ALIGN_PARAGRAPH.LEFT)):
            s = st[name]
            _set_font(s, size=size, bold=True, italic=False, caps=caps)
            s.paragraph_format.alignment = align
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.line_spacing = 1.5
            s.paragraph_format.space_before = Pt(0 if name == "Heading 1" else 12)
            s.paragraph_format.space_after = Pt(6)
            if name != "Heading 1":
                s.paragraph_format.first_line_indent = Cm(1.25)
        st["Heading 1"].paragraph_format.page_break_before = True
        for lvl in (1, 2):
            s = st.add_style(f"toc {lvl}", WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st["Normal"]
            _set_font(s, size=14)
            s.paragraph_format.line_spacing = 1.15
            s.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.8)
            s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W["A4"]), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        sec = d.sections[0]
        self._geom(sec, "A4")
        return d

    def _geom(self, sec, kind):
        w, h = PAGE[kind]
        sec.orientation = WD_ORIENT.LANDSCAPE if kind in ("A4L", "A3L") else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.header_distance = Cm(0.8)
        sec.footer_distance = Cm(0.8)

    def _para(self, d, text="", align=None, size=None, bold=None, indent=False, spacing=None, keep=False, style=None):
        p = d.add_paragraph(style=style)
        pf = p.paragraph_format
        if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else: p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(1.25) if indent else Cm(0)
        if spacing is not None: pf.line_spacing = spacing
        if keep: pf.keep_with_next = True
        if text: self._runs(p, text, size=size, bold=bold)
        return p

    def _runs(self, p, text, size=None, bold=None, italic=None):
        text = self._res(text)
        for part in INLINE.split(text):
            if not part: continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                r = p.add_run(part[2:-2]); _run_font(r, size=size, bold=True, italic=italic)
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                r = p.add_run(part[1:-1]); _run_font(r, name=MONO, size=(size or 14) - 2 if (size or 14) > 10 else size)
                if bold: r.font.bold = True
            elif part.startswith("__") and part.endswith("__") and len(part) > 4:
                r = p.add_run(part[2:-2]); _run_font(r, size=size, bold=bold, italic=True)
            else:
                r = p.add_run(part); _run_font(r, size=size, bold=bold, italic=italic)

    # --- титульные элементы
    def _approval_sheet(self, d, total):
        _p = self._para
        _p(d, f"{self.designation}-ЛУ", align="c", size=12, spacing=1.0)
        _p(d, "", spacing=1.0)
        t = d.add_table(rows=1, cols=2); t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t.rows[0].cells
        c0.width, c1.width = Cm(8.0), Cm(9.0)
        c0.text = ""
        c1.text = ""
        for ln in (("УТВЕРЖДАЮ", True), (self.org, False), ("", False), ("______________ / Силантьев Н. /", False), ("«____» ______________ " + self.year + " г.", False)):
            pp = c1.add_paragraph() if c1.paragraphs[0].text or c1.paragraphs[0].runs else c1.paragraphs[0]
            pp.paragraph_format.line_spacing = 1.15
            r = pp.add_run(ln[0]); _run_font(r, size=13, bold=ln[1])
        c0.paragraphs[0].add_run("")
        for _ in range(7): _p(d, "", spacing=1.0)
        _p(d, self.program.upper(), align="c", bold=True, size=16, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, self.title.upper(), align="c", bold=True, size=16, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, "ЛИСТ УТВЕРЖДЕНИЯ", align="c", bold=True, size=16, spacing=1.15)
        _p(d, self.designation, align="c", size=14, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, f"Листов {total}", align="c", size=14, spacing=1.15)
        for _ in range(5): _p(d, "", spacing=1.0)
        t = d.add_table(rows=2, cols=5); t.autofit = False
        _borders(t)
        for i, h in enumerate(("Инв. № подл.", "Подп. и дата", "Взам. инв. №", "Инв. № дубл.", "Подп. и дата")):
            c = t.rows[0].cells[i]; c.width = Cm(3.4)
            c.text = ""; r = c.paragraphs[0].add_run(h); _run_font(r, size=10)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            t.rows[1].cells[i].width = Cm(3.4)
        t.rows[1].height = Cm(1.4)

    def _title_page(self, d, total):
        _p = self._para
        _p(d, "", spacing=1.0)
        for _ in range(8): _p(d, "", spacing=1.0)
        _p(d, self.program.upper(), align="c", bold=True, size=18, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, self.title.upper(), align="c", bold=True, size=18, spacing=1.15)
        if self.subtitle:
            _p(d, "", spacing=1.0)
            _p(d, self.subtitle, align="c", size=14, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, self.designation, align="c", size=14, spacing=1.15)
        _p(d, "", spacing=1.0)
        _p(d, f"Листов {total}", align="c", size=14, spacing=1.15)
        for _ in range(9): _p(d, "", spacing=1.0)
        _p(d, self.year, align="c", size=14, spacing=1.15)

    def _toc(self, d, pages):
        h = self._para(d, "СОДЕРЖАНИЕ", align="c", bold=True, spacing=1.5)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.space_after = Pt(6)
        n = len(self.toc_entries)
        for i, (lvl, label) in enumerate(self.toc_entries):
            p = d.add_paragraph(style=f"toc {lvl}")
            if i == 0:
                for t, ins in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
                    r = p.add_run(); _run_font(r)
                    if t:
                        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
                    else:
                        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = f" {ins} "; r._r.append(e)
            r = p.add_run(label.replace("\n", " ")); _run_font(r, size=14)
            r = p.add_run("\t" + str(pages.get(i, ""))); _run_font(r, size=14)
            if i == n - 1:
                r = p.add_run(); _run_font(r)
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)

    def _header(self, sec):
        sec.header.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.line_spacing = 1.0
        _fld(hp, "PAGE", size=12)
        p2 = sec.header.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.line_spacing = 1.0
        r = p2.add_run(self.designation); _run_font(r, size=12)

    def _sect(self, d, kind, first=False):
        sec = d.add_section(WD_SECTION.NEW_PAGE)
        self._geom(sec, kind)
        sec.different_first_page_header_footer = False
        pg = sec._sectPr.find(qn("w:pgNumType"))
        if pg is not None and not first: sec._sectPr.remove(pg)
        sec.header.is_linked_to_previous = True
        return sec

    # --- блоки
    def _render_block(self, d, b, cur):
        t = b["t"]
        if t == "h":
            st = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[b["lvl"]]
            p = d.add_paragraph(style=st)
            self._runs(p, b["label"], bold=True)
            if b["lvl"] == 1 and cur in ("A3L", "A4L"):
                p.paragraph_format.page_break_before = False
        elif t == "app":
            p = d.add_paragraph(style="Heading 1")
            r = p.add_run(f"ПРИЛОЖЕНИЕ {b['n']}"); _run_font(r, bold=True)
            r.add_break()
            r = p.add_run(f"({b['status']})"); _run_font(r, bold=False)
            r.add_break()
            r = p.add_run(b["title"]); _run_font(r, bold=True)
        elif t == "p":
            p = self._para(d, b["text"], align=b["align"], indent=b["indent"], bold=b["bold"] or None, size=b["size"], keep=b["keep"])
        elif t == "li":
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent = Cm(2.1); pf.first_line_indent = Cm(-0.85)
            pf.tab_stops.add_tab_stop(Cm(2.1))
            r = p.add_run(b["mark"] + "\t"); _run_font(r)
            self._runs(p, b["text"])
        elif t == "note":
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.25)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
            r = p.add_run(b["label"] + " – "); _run_font(r, size=12)
            self._runs(p, b["text"], size=12)
        elif t == "code":
            lines = b["text"].split("\n")
            for i, ln in enumerate(lines):
                p = d.add_paragraph()
                pf = p.paragraph_format
                pf.line_spacing = 1.0; pf.left_indent = Cm(0.3)
                pf.keep_together = True
                r = p.add_run(ln if ln else " "); _run_font(r, name=MONO, size=b["size"])
                pPr = p._p.get_or_add_pPr()
                sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "F2F2F2"); pPr.append(sh)
            self._para(d, "", spacing=0.6)
        elif t == "sp":
            p = self._para(d, "", spacing=1.0); p.paragraph_format.space_after = Pt(b["pts"])
        elif t == "pb":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "fig":
            tw = TEXT_W[cur]
            maxh = {"A4": 20.5, "A3L": 22.0, "A3P": 35.5, "A4L": 14.0}[cur]
            sc = min(tw / b["nw"], maxh / b["nh"])
            wcm = b["nw"] * sc
            p = self._para(d, "", align="c", spacing=1.0, keep=True)
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(); r.add_picture(b["path"], width=Cm(wcm))
            c = self._para(d, f"Рисунок {b['n']} – {self._res(b['caption'])}", align="c", size=14, spacing=1.15)
            c.paragraph_format.space_after = Pt(10)
        elif t == "tab":
            self._table(d, b, cur)

    def _table(self, d, b, cur):
        tw = TEXT_W[cur]
        if b["n"] is not None:
            c = self._para(d, f"Таблица {b['n']} – {self._res(b['caption'])}", align="l", size=14, spacing=1.15, keep=True)
            c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(3)
        ncol = len(b["header"])
        wts = b["widths"] or [1] * ncol
        tot = float(sum(wts))
        ws = [tw * w / tot for w in wts]
        t = d.add_table(rows=1, cols=ncol)
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        _borders(t); _cell_margins(t)
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        tw_el = tblPr.find(qn("w:tblW"))
        if tw_el is None:
            tw_el = OxmlElement("w:tblW"); tblPr.append(tw_el)
        tw_el.set(qn("w:type"), "dxa"); tw_el.set(qn("w:w"), str(int(tw * 567)))
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(ws[i] * 567)))
        f = b["font"]

        def fill(cell, text, bold=False, align="l"):
            cell.width = Cm(ws[cell_idx[0]])
            lines = self._res(str(text)).split("\n")
            for i, ln in enumerate(lines):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER}[align]
                self._runs(p, ln, size=f, bold=bold or None)

        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        cell_idx = [0]
        for i, h in enumerate(b["header"]):
            cell_idx[0] = i
            fill(hdr.cells[i], h, bold=True, align="c"); _shade(hdr.cells[i], "D9D9D9")
        for row in b["rows"]:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cell_idx[0] = i
                fill(cells[i], v)
        self._para(d, "", spacing=0.8)

    def render(self, path, pages=None, total=0):
        pages = pages or {}
        d = self._new_doc()
        if self.kind == "espd":
            self._approval_sheet(d, total)
            sec = self._sect(d, "A4", first=True)
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); sec._sectPr.append(pg)
            sec.different_first_page_header_footer = True
            self._header(sec)
            self._title_page(d, total)
            if self.want_toc and self.toc_entries:
                self._toc(d, pages)
        else:
            sec = d.sections[0]
            sec.different_first_page_header_footer = False
            self._header(sec)
            self._legal_title(d)
        cur = "A4"
        base = None
        base_at = []  # режим по умолчанию на момент каждого блока
        _b = None
        for b in self.blocks:
            if b["t"] == "mode": _b = b["sect"]
            base_at.append(_b)
        for i, b in enumerate(self.blocks):
            if b["t"] == "mode":
                continue
            need = b.get("sect") or base_at[i] or "A4"
            if b["t"] in ("h", "app"):  # заголовок остаётся на одной странице со своим содержимым
                for j in range(i + 1, len(self.blocks)):
                    nb = self.blocks[j]
                    if nb["t"] not in ("h", "app", "mode"):
                        need = nb.get("sect") or base_at[j] or "A4"; break
            if need != cur:
                self._sect(d, need)
                cur = need
            self._render_block(d, b, cur)
        if cur != "A4":
            self._sect(d, "A4"); cur = "A4"
        if self.kind == "espd":
            self._change_sheet(d)
        d.core_properties.title = f"{self.program}. {self.title}"
        d.core_properties.author = "NEXO"
        d.core_properties.language = "ru-RU"
        d.save(path)

    def _legal_title(self, d):
        p = self._para(d, self.title.upper(), align="c", bold=True, size=16, spacing=1.15)
        p.paragraph_format.space_after = Pt(6)
        if self.subtitle:
            self._para(d, self.subtitle, align="c", size=13, spacing=1.15)
        self._para(d, f"{self.designation} · версия {self.version}", align="c", size=12, spacing=1.15).paragraph_format.space_after = Pt(12)

    def _change_sheet(self, d):
        h = d.add_paragraph(style="Heading 1")
        r = h.add_run("ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ"); _run_font(r, bold=True)
        cols = ["Изм.", "Номера листов (страниц)\nизменённых", "заменённых", "новых", "аннулированных",
                "Всего листов (страниц)\nв докум.", "№ докум.", "Входящий № сопроводит.\nдокум. и дата", "Подп.", "Дата"]
        t = d.add_table(rows=2, cols=len(cols)); t.autofit = False
        _borders(t); _cell_margins(t)
        w = [1.0, 1.8, 1.8, 1.5, 2.0, 1.9, 1.4, 2.6, 1.5, 1.5]
        s = sum(w); w = [17.0 * x / s for x in w]
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]; cell.width = Cm(w[i]); cell.text = ""
            r = cell.paragraphs[0].add_run(c); _run_font(r, size=9, bold=True)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            _shade(cell, "D9D9D9")
        for _ in range(12):
            cells = t.add_row().cells
            for i in range(len(cols)): cells[i].width = Cm(w[i])
        t.rows[1]._tr.getparent().remove(t.rows[1]._tr)

    # ------------------------------------------------------------ двухпроходная сборка
    def build(self, outdir, name):
        os.makedirs(outdir, exist_ok=True)
        out = os.path.join(outdir, name)
        pages, total = {}, 0
        for it in range(3):
            self.render(out, pages, total)
            newpages, newtotal = self._measure(out)
            if newpages == pages and newtotal == total: break
            pages, total = newpages, newtotal
        self.render(out, pages, total)
        if self.refkey:
            import json
            json.dump(self.keys, open(os.path.join(self.ext_dir, f"refs_{self.refkey}.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=0)
        return out, total

    def _measure(self, docx_path):
        tmp = tempfile.mkdtemp()
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, docx_path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        pdf = os.path.join(tmp, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
        txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
        pgs = txt.split("\f")
        if pgs and not pgs[-1].strip(): pgs = pgs[:-1]
        n_pdf = len(pgs)
        total = n_pdf - 1 if self.kind == "espd" else n_pdf  # лист утверждения не учитывается
        norm = lambda s: re.sub(r"\s+", " ", s).strip().casefold()
        # начало поиска — после страницы(страниц) содержания
        start = 0
        if self.want_toc:
            for i, pg in enumerate(pgs):
                if "содержание" in norm(pg)[:200] and i >= 1:
                    start = i; break
            # содержание может занимать несколько страниц: пропускаем страницы с точечными выносками
            while start + 1 < n_pdf and "....." in pgs[start + 1][:3000] and pgs[start + 1].count(".....") > 5:
                start += 1
            start += 1
        pages = {}
        ptr = start
        for idx, (lvl, label) in enumerate(self.toc_entries):
            key = norm(label.replace("\n", " "))[:28]
            if label.startswith("ПРИЛОЖЕНИЕ"):
                key = norm(" ".join(label.split()[:2]))
            found = None
            for pi in range(ptr, n_pdf):
                for ln in pgs[pi].split("\n"):
                    if norm(ln).startswith(key):
                        found = pi; break
                if found is not None: break
            if found is None:
                print(f"  [!] заголовок не найден в PDF: {label[:60]}")
                continue
            pages[idx] = found + 1 - (1 if self.kind == "espd" else 0)
            ptr = found
        shutil.rmtree(tmp, ignore_errors=True)
        return pages, total
